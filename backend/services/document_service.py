"""Document Import/Export and Subtitle Generation Service for LinguaFusion.

Supports:
- Importing text from .txt, .docx (python-docx), and .pdf (PyMuPDF / fitz).
- Exporting structured .txt, .docx, and styled .pdf documents.
- Exporting speech transcription segments as .srt and .vtt subtitle files.
"""

import io
import os
from pathlib import Path
from typing import Any, Sequence

# Dependencies checked in .venv: docx, fitz (PyMuPDF), reportlab
try:
    import docx
except ImportError:
    docx = None

try:
    import fitz  # PyMuPDF
except ImportError:
    fitz = None

try:
    from reportlab.lib.pagesizes import letter
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer
except ImportError:
    letter = None
    SimpleDocTemplate = None


def extract_text_from_file(file_path: str | Path) -> str:
    """Extract plain text from a .txt, .docx, or .pdf file."""
    path = Path(file_path)
    if not path.exists():
        raise FileNotFoundError(f"File not found: {file_path}")

    ext = path.suffix.lower()

    if ext in (".txt", ".md"):
        for encoding in ("utf-8", "utf-8-sig", "latin-1", "cp1252"):
            try:
                return path.read_text(encoding=encoding)
            except UnicodeDecodeError:
                continue
        return path.read_text(errors="ignore")

    elif ext == ".docx":
        if docx is None:
            raise RuntimeError("python-docx is not installed in the environment.")
        doc = docx.Document(str(path))
        full_text = [p.text for p in doc.paragraphs if p.text.strip()]
        return "\n\n".join(full_text)

    elif ext == ".pdf":
        if fitz is None:
            raise RuntimeError("PyMuPDF (fitz) is not installed in the environment.")
        doc = fitz.open(str(path))
        pages_text = []
        for page in doc:
            text = page.get_text("text")
            if text.strip():
                pages_text.append(text.strip())
        doc.close()
        return "\n\n".join(pages_text)

    else:
        raise ValueError(f"Unsupported file format: '{ext}'. Supported formats are .txt, .md, .docx, .pdf")


def export_document(text: str, format_type: str = "txt", title: str = "LinguaFusion Document") -> bytes:
    """Export text as a .txt, .md, .docx, or .pdf file returned as raw bytes."""
    fmt = format_type.lower().strip(".")

    if fmt in ("txt", "md"):
        return text.encode("utf-8")

    elif fmt == "docx":
        if docx is None:
            raise RuntimeError("python-docx is not installed.")
        doc = docx.Document()
        doc.add_heading(title, level=1)
        paragraphs = text.split("\n\n")
        for p_text in paragraphs:
            clean_p = p_text.strip()
            if clean_p:
                doc.add_paragraph(clean_p)
        buffer = io.BytesIO()
        doc.save(buffer)
        return buffer.getvalue()

    elif fmt == "pdf":
        if SimpleDocTemplate is None:
            raise RuntimeError("reportlab is not installed.")
        buffer = io.BytesIO()
        doc = SimpleDocTemplate(
            buffer,
            pagesize=letter,
            rightMargin=54,
            leftMargin=54,
            topMargin=54,
            bottomMargin=54,
        )
        styles = getSampleStyleSheet()
        title_style = ParagraphStyle(
            "DocTitle",
            parent=styles["Title"],
            fontSize=20,
            leading=24,
            textColor="#1A202C",
            spaceAfter=20,
        )
        body_style = ParagraphStyle(
            "DocBody",
            parent=styles["Normal"],
            fontSize=11,
            leading=16,
            textColor="#2D3748",
            spaceAfter=12,
        )

        story = [Paragraph(title, title_style), Spacer(1, 12)]
        paragraphs = text.split("\n\n")
        for p_text in paragraphs:
            clean_p = p_text.strip().replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
            if clean_p:
                # Replace newlines inside paragraph with line breaks
                clean_p = clean_p.replace("\n", "<br/>")
                story.append(Paragraph(clean_p, body_style))

        doc.build(story)
        return buffer.getvalue()

    else:
        raise ValueError(f"Unsupported export format: '{fmt}'. Use 'txt', 'docx', or 'pdf'.")


def format_srt_timestamp(seconds: float) -> str:
    """Format seconds into SRT timestamp HH:MM:SS,mmm."""
    seconds = max(0.0, seconds)
    hrs = int(seconds // 3600)
    mins = int((seconds % 3600) // 60)
    secs = int(seconds % 60)
    millis = int(round((seconds - int(seconds)) * 1000))
    if millis >= 1000:
        secs += 1
        millis = 0
    return f"{hrs:02d}:{mins:02d}:{secs:02d},{millis:03d}"


def format_vtt_timestamp(seconds: float) -> str:
    """Format seconds into WebVTT timestamp HH:MM:SS.mmm."""
    seconds = max(0.0, seconds)
    hrs = int(seconds // 3600)
    mins = int((seconds % 3600) // 60)
    secs = int(seconds % 60)
    millis = int(round((seconds - int(seconds)) * 1000))
    if millis >= 1000:
        secs += 1
        millis = 0
    return f"{hrs:02d}:{mins:02d}:{secs:02d}.{millis:03d}"


def export_subtitles(segments: Sequence[dict[str, Any]], format_type: str = "srt") -> str:
    """Export list of speech segments ({'start': float, 'end': float, 'text': str}) as SRT or VTT string."""
    fmt = format_type.lower().strip(".")
    lines = []

    if fmt == "vtt":
        lines.append("WEBVTT")
        lines.append("")
        for idx, seg in enumerate(segments, 1):
            start_str = format_vtt_timestamp(float(seg.get("start", 0.0)))
            end_str = format_vtt_timestamp(float(seg.get("end", 0.0)))
            text = str(seg.get("text", "")).strip()
            lines.append(f"{idx}")
            lines.append(f"{start_str} --> {end_str}")
            lines.append(text)
            lines.append("")
    else:  # srt
        for idx, seg in enumerate(segments, 1):
            start_str = format_srt_timestamp(float(seg.get("start", 0.0)))
            end_str = format_srt_timestamp(float(seg.get("end", 0.0)))
            text = str(seg.get("text", "")).strip()
            lines.append(f"{idx}")
            lines.append(f"{start_str} --> {end_str}")
            lines.append(text)
            lines.append("")

    return "\n".join(lines)


def batch_process_documents(files_data: list, source_lang: str, target_lang: str, translate_func, export_format: str = "same") -> bytes:
    """Takes a list of (filename, file_bytes) tuples, translates each document, formats output, and returns a zip file containing translated documents."""
    import zipfile
    import tempfile

    zip_buffer = io.BytesIO()
    with zipfile.ZipFile(zip_buffer, "w", zipfile.ZIP_DEFLATED) as zip_file:
        for filename, file_bytes in files_data:
            path_obj = Path(filename)
            ext = path_obj.suffix.lower()
            stem = path_obj.stem

            with tempfile.NamedTemporaryFile(suffix=ext, delete=False) as tmp:
                tmp.write(file_bytes)
                tmp_path = tmp.name

            try:
                raw_text = extract_text_from_file(tmp_path)
            finally:
                try:
                    os.remove(tmp_path)
                except Exception:
                    pass

            if not raw_text.strip():
                continue

            translated_text = translate_func(raw_text, source_lang, target_lang)
            if export_format in ("docx", "pdf", "txt", "md"):
                out_fmt = export_format
            else:
                out_fmt = "docx" if ext == ".docx" else ("pdf" if ext == ".pdf" else ("md" if ext == ".md" else "txt"))

            out_bytes = export_document(translated_text, format_type=out_fmt, title=f"Translated - {stem}")
            out_filename = f"{stem}_translated_{target_lang}.{out_fmt}"
            zip_file.writestr(out_filename, out_bytes)

    return zip_buffer.getvalue()


def batch_process_ocr(files_data: list, ocr_lang: str, target_lang: str, ocr_func, translate_func, export_format: str = "docx") -> bytes:
    """Takes a list of (filename, image/pdf bytes), extracts text using OCR, optionally translates, formats output into choosable format, and returns zip bytes."""
    import zipfile
    import tempfile

    zip_buffer = io.BytesIO()
    with zipfile.ZipFile(zip_buffer, "w", zipfile.ZIP_DEFLATED) as zip_file:
        for filename, file_bytes in files_data:
            path_obj = Path(filename)
            ext = path_obj.suffix.lower()
            stem = path_obj.stem

            with tempfile.NamedTemporaryFile(suffix=ext, delete=False) as tmp:
                tmp.write(file_bytes)
                tmp_path = tmp.name

            try:
                extracted_text = ocr_func(tmp_path, ocr_lang)
            finally:
                try:
                    os.remove(tmp_path)
                except Exception:
                    pass

            if not extracted_text.strip():
                continue

            final_text = extracted_text
            if target_lang and target_lang != "none":
                translated = translate_func(extracted_text, ocr_lang, target_lang)
                final_text = f"Extracted Text:\n{extracted_text}\n\nTranslation ({target_lang}):\n{translated}"

            out_fmt = export_format if export_format in ("docx", "pdf", "txt") else "docx"
            out_bytes = export_document(final_text, format_type=out_fmt, title=f"OCR - {stem}")
            out_filename = f"{stem}_ocr.{out_fmt}"
            zip_file.writestr(out_filename, out_bytes)

    return zip_buffer.getvalue()
